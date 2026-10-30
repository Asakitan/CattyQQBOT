from __future__ import annotations

import asyncio
from contextlib import closing
import importlib
import importlib.util
import io
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
import types
import unittest
from unittest.mock import patch
import zipfile


PACKAGE = "_catty_knowledge_base_tests"
if PACKAGE not in sys.modules:
    package = types.ModuleType(PACKAGE)
    package.__path__ = [str(Path(__file__).resolve().parents[1] / "src" / "catty_qq_ai")]
    sys.modules[PACKAGE] = package
MODULE = importlib.import_module(PACKAGE + ".knowledge_base")
TaskIdentity = importlib.import_module(PACKAGE + ".task_contracts").TaskIdentity
KnowledgeBase = MODULE.KnowledgeBase
KnowledgeBaseError = MODULE.KnowledgeBaseError


class FakeEmbedding:
    def __init__(self):
        self.available = True
        self.calls = []

    @staticmethod
    def name():
        return "fake-semantic-space-v1"

    def __call__(self, texts):
        self.calls.extend(texts)
        if not self.available:
            raise RuntimeError("not ready")
        return [[1.0, 0.0, 0.1] if any(term in text for term in ("car", "automobile", "vehicle"))
                else [0.0, 1.0, 0.1] for text in texts]

    def embed_query(self, texts):
        return self(texts)


class KnowledgeBaseTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name) / "knowledge"
        self.identity = TaskIdentity("100", "group:200", "fadianji")
        self.settings = {"parse_in_worker": False, "chunk_size": 80, "chunk_overlap": 10}
        self.store = KnowledgeBase(self.root, self.settings)
        self.addCleanup(self.store.close)

    def test_chinese_fts_and_citation_round_trip(self):
        document = self.store.ingest_text(self.identity, "manual.md", "# 数据库维护\n每天进行数据库备份。\f# 恢复\n恢复操作必须先校验备份。", "https://example.test/manual")
        hits = self.store.search(self.identity, "数据库备份")
        self.assertTrue(hits)
        self.assertIn("数据库备份", hits[0]["text"])
        self.assertEqual(hits[0]["search_mode"], "fts")
        self.assertTrue(hits[0]["semantic_pending"])
        text = self.store.read_artifact(self.identity, document["text_artifact_id"]).decode("utf-8")
        for hit in hits:
            self.assertEqual(text[hit["char_start"]:hit["char_end"]], hit["text"])
            self.assertEqual(self.store.get_chunk(self.identity, hit["chunk_id"])["citation"], hit["citation"])
            self.assertEqual(hit["citation"]["source"], "https://example.test/manual")
        page_hit = self.store.search(self.identity, "恢复")[0]
        self.assertEqual(page_hit["page"], 2)
        self.assertEqual(page_hit["title"], "恢复")

    def test_persistence_deduplication_and_versions(self):
        first = self.store.ingest_text(self.identity, "notes.txt", "uniqueold keyword")
        duplicate = self.store.ingest_text(self.identity, "copy.txt", "uniqueold keyword")
        self.assertTrue(duplicate["duplicate"])
        self.assertEqual(first["doc_id"], duplicate["doc_id"])
        updated = self.store.ingest_text(self.identity, "notes.txt", "uniquenew keyword")
        self.assertEqual(updated["version"], 2)
        self.assertEqual(updated["doc_id"], first["doc_id"])
        self.assertIsNone(self.store.read_artifact(self.identity, first["artifact_id"]))
        self.assertFalse((self.root / "artifacts" / first["artifact_id"]).exists())
        reopened = KnowledgeBase(self.root, self.settings)
        self.addCleanup(reopened.close)
        self.assertEqual(reopened.search(self.identity, "uniqueold"), [])
        self.assertEqual(reopened.search(self.identity, "uniquenew")[0]["version"], 2)

    def test_identity_isolation_includes_user_scope_and_persona_even_for_owner(self):
        document = self.store.ingest_text(self.identity, "private.txt", "confidentialtoken")
        chunk_id = self.store.search(self.identity, "confidentialtoken")[0]["chunk_id"]
        others = [TaskIdentity("101", "group:200", "fadianji"), TaskIdentity("100", "private:100", "fadianji"),
                  TaskIdentity("100", "group:200", "catty"), TaskIdentity("999", "group:200", "fadianji", True)]
        for other in others:
            with self.subTest(identity=other):
                self.assertEqual(self.store.list_documents(other), [])
                self.assertEqual(self.store.search(other, "confidentialtoken"), [])
                self.assertIsNone(self.store.get_chunk(other, chunk_id))
                self.assertIsNone(self.store.read_artifact(other, document["artifact_id"]))
                self.assertFalse(self.store.delete_document(other, document["doc_id"])["ok"])
                self.assertFalse(self.store.rebuild(other, document["doc_id"])["ok"])
                independent = self.store.ingest_text(other, "private.txt", "confidentialtoken")
                self.assertNotEqual(independent["doc_id"], document["doc_id"])
                self.assertNotEqual(independent["artifact_id"], document["artifact_id"])
        with self.assertRaises(TypeError):
            self.store.search({"user_id": "100", "scope": "group:200", "persona": "fadianji"}, "confidentialtoken")

    def test_explicit_admin_listing_contains_metadata_not_document_content(self):
        other = TaskIdentity("101", "private:101", "catty")
        self.store.ingest_text(self.identity, "first.txt", "secret_body_one")
        self.store.ingest_text(other, "second.txt", "secret_body_two")
        rows = self.store.admin_list_documents()
        self.assertEqual({(row["user_id"], row["scope"], row["persona"]) for row in rows},
                         {("100", "group:200", "fadianji"), ("101", "private:101", "catty")})
        for row in rows:
            self.assertNotIn("text", row)
            self.assertNotIn("content", row)
            self.assertNotIn("artifact_id", row)
            self.assertNotIn("text_artifact_id", row)
        self.assertEqual(len(self.store.list_documents(self.identity)), 1)
        self.assertEqual(self.store.search(self.identity, "secret_body_two"), [])

    def test_hybrid_semantics_real_fts_and_diversity(self):
        embedding = FakeEmbedding()
        store = KnowledgeBase(self.root / "vectors", self.settings, embedding)
        self.addCleanup(store.close)
        direct = store.ingest_text(self.identity, "car.txt", "car vehicle maintenance. " * 12)
        synonym = store.ingest_text(self.identity, "synonym.txt", "automobile service intervals")
        store.ingest_text(self.identity, "unrelated.txt", "garden plants flowers")
        hits = store.search(self.identity, "car", 3)
        self.assertEqual(hits[0]["doc_id"], direct["doc_id"])
        self.assertIsNotNone(hits[0]["lexical_rank"])
        self.assertIsNotNone(hits[0]["semantic_score"])
        self.assertEqual(hits[0]["search_mode"], "hybrid")
        self.assertIn(synonym["doc_id"], {hit["doc_id"] for hit in hits})
        self.assertTrue(all(not hit["semantic_pending"] and hit["citation"]["artifact_id"] for hit in hits))
        self.assertLessEqual(sum(hit["doc_id"] == direct["doc_id"] for hit in hits), 2)
        reopened = KnowledgeBase(self.root / "vectors", self.settings, FakeEmbedding())
        self.addCleanup(reopened.close)
        self.assertTrue(any(hit["doc_id"] == synonym["doc_id"] for hit in reopened.search(self.identity, "vehicle", 3)))

    def test_missing_bge_is_honest_and_retry_backfills_vectors(self):
        backend = importlib.import_module(PACKAGE + ".rag_embedding_backend")
        unavailable = backend.LocalBGEEmbeddingFunction(model_name="")
        self.addCleanup(unavailable.close)
        store = KnowledgeBase(self.root / "pending", self.settings, unavailable)
        self.addCleanup(store.close)
        document = store.ingest_text(self.identity, "car.txt", "automobile vehicle car")
        self.assertTrue(document["semantic_pending"])
        self.assertEqual(document["semantic_error"], "LocalBGENotReadyError")
        self.assertEqual(store.search(self.identity, "car")[0]["search_mode"], "fts")
        with closing(sqlite3.connect(store.database_path)) as database, database:
            self.assertEqual(database.execute("SELECT COUNT(*) FROM vectors").fetchone()[0], 0)
        recovered = KnowledgeBase(self.root / "pending", self.settings, FakeEmbedding())
        self.addCleanup(recovered.close)
        retry = recovered.retry_pending(self.identity)
        self.assertEqual(retry, {"ok": True, "indexed": 1, "failed": 0, "pending": 0})
        self.assertFalse(recovered.list_documents(self.identity)[0]["semantic_pending"])

    def test_invalid_vectors_never_become_ready(self):
        for vector in ([0.0, 0.0], [float("nan"), 1.0], [True, 1.0]):
            with self.subTest(vector=vector):
                store = KnowledgeBase(self.root / str(len(list(self.root.iterdir()))), self.settings, lambda texts: [vector for text in texts])
                self.addCleanup(store.close)
                result = store.ingest_text(self.identity, "data.txt", "searchable text")
                self.assertTrue(result["semantic_pending"])
                self.assertEqual(result["semantic_error"], "invalid_embeddings")

    def test_delete_and_rebuild_touch_only_stored_artifacts(self):
        embedding = FakeEmbedding()
        store = KnowledgeBase(self.root / "delete", self.settings, embedding)
        self.addCleanup(store.close)
        document = store.ingest_text(self.identity, "car.txt", "automobile car")
        chunk = store.search(self.identity, "car")[0]
        with closing(sqlite3.connect(store.database_path)) as database, database:
            database.execute("DELETE FROM chunks_fts")
            database.execute("DELETE FROM vectors")
        rebuilt = store.rebuild(self.identity, document["doc_id"])
        self.assertTrue(rebuilt["rebuilt"])
        self.assertEqual(rebuilt["version"], 1)
        self.assertEqual(store.search(self.identity, "car")[0]["chunk_id"], chunk["chunk_id"])
        self.assertTrue(store.delete_document(self.identity, document["doc_id"])["ok"])
        self.assertEqual(store.list_documents(self.identity), [])
        self.assertEqual(store.search(self.identity, "car"), [])
        self.assertIsNone(store.get_chunk(self.identity, chunk["chunk_id"]))
        self.assertIsNone(store.read_artifact(self.identity, rebuilt["artifact_id"]))
        self.assertEqual(list((store.root / "artifacts").iterdir()), [])
        with closing(sqlite3.connect(store.database_path)) as database, database:
            for table in ("documents", "chunks", "chunks_fts", "vectors", "artifacts", "retired_artifacts"):
                self.assertEqual(database.execute("SELECT COUNT(*) FROM " + table).fetchone()[0], 0)
        self.assertFalse(store.delete_document(self.identity, document["doc_id"])["ok"])

    def test_cleanup_failure_is_honest_persistent_and_retryable(self):
        original = self.store.ingest_text(self.identity, "notes.txt", "oldtext")
        with patch.object(Path, "unlink", side_effect=PermissionError("locked")):
            updated = self.store.ingest_text(self.identity, "notes.txt", "newtext")
            result = self.store.delete_document(self.identity, updated["doc_id"])
        self.assertTrue(updated["cleanup_pending"])
        self.assertFalse(result["ok"])
        self.assertTrue(result["deleted"])
        self.assertTrue(result["cleanup_pending"])
        self.assertEqual(self.store.search(self.identity, "newtext"), [])
        self.assertIsNone(self.store.read_artifact(self.identity, original["artifact_id"]))
        self.assertEqual(len(list((self.root / "artifacts").iterdir())), 4)
        reopened = KnowledgeBase(self.root, self.settings)
        self.addCleanup(reopened.close)
        other = TaskIdentity("999", self.identity.scope, self.identity.persona, True)
        self.assertFalse(reopened.delete_document(other, original["doc_id"])["ok"])
        self.assertTrue(reopened.delete_document(self.identity, original["doc_id"])["ok"])
        self.assertEqual(list((self.root / "artifacts").iterdir()), [])

    def test_settings_instances_cannot_bypass_hard_upload_limit(self):
        settings = MODULE.KnowledgeBaseSettings(max_upload_bytes=20 * 1024 * 1024 + 1)
        with self.assertRaises(KnowledgeBaseError) as caught:
            KnowledgeBase(self.root / "oversized", settings)
        self.assertEqual(caught.exception.code, "invalid_settings")

    def test_html_does_not_index_scripts_or_fetch_links(self):
        self.store.ingest_text(self.identity, "page.html", '<title>Manual</title><script>secrettoken</script><p>publictoken &amp; text</p><img src="file:///private">')
        self.assertEqual(self.store.search(self.identity, "secrettoken"), [])
        self.assertEqual(self.store.search(self.identity, "publictoken")[0]["title"], "Manual")

    def test_rejects_paths_limits_empty_text_and_empty_html(self):
        for name in ("../a.txt", "..\\a.txt", "C:\\a.txt", "/a.txt", "file.txt:stream", "bad\x00.txt", "bad.exe"):
            with self.subTest(name=name), self.assertRaises(KnowledgeBaseError):
                self.store.ingest_text(self.identity, name, "text")
        for name, text in (("blank.txt", " \n"), ("blank.html", "<script>hidden</script>")):
            with self.subTest(name=name), self.assertRaises(KnowledgeBaseError) as caught:
                self.store.ingest_text(self.identity, name, text)
            self.assertEqual(caught.exception.code, "no_text")
        limited = KnowledgeBase(self.root / "limited", {**self.settings, "max_upload_bytes": 10})
        self.addCleanup(limited.close)
        with self.assertRaises(KnowledgeBaseError) as caught:
            limited.ingest_bytes(self.identity, "huge.txt", b"x" * 11)
        self.assertEqual(caught.exception.code, "upload_too_large")
        self.assertEqual(limited.list_documents(self.identity), [])
        self.assertIsNone(self.store.read_artifact(self.identity, "../../config.json"))

    def test_zip_bomb_traversal_and_entry_count_are_rejected_before_parser(self):
        fixtures = [("word/document.xml", b"x" * 100000, "archive_ratio_exceeded"),
                    ("../outside.xml", b"text", "unsafe_archive")]
        for name, data, error in fixtures:
            output = io.BytesIO()
            with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr(name, data)
            with self.subTest(name=name), self.assertRaises(KnowledgeBaseError) as caught:
                self.store.ingest_bytes(self.identity, "malicious.docx", output.getvalue())
            self.assertEqual(caught.exception.code, error)
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("one.xml", "text")
            archive.writestr("two.xml", "text")
        limited = KnowledgeBase(self.root / "entries", {**self.settings, "max_archive_entries": 1})
        self.addCleanup(limited.close)
        with self.assertRaises(KnowledgeBaseError) as caught:
            limited.ingest_bytes(self.identity, "malicious.xlsx", output.getvalue())
        self.assertEqual(caught.exception.code, "archive_entries_exceeded")

    def test_archive_total_size_and_xml_entities_are_bounded(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("one.xml", "12345")
            archive.writestr("two.xml", "12345")
        limited = KnowledgeBase(self.root / "archive", {**self.settings, "max_archive_bytes": 8})
        self.addCleanup(limited.close)
        with self.assertRaises(KnowledgeBaseError) as caught:
            limited.ingest_bytes(self.identity, "large.docx", output.getvalue())
        self.assertEqual(caught.exception.code, "archive_size_exceeded")
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("word/document.xml", '<!DOCTYPE doc [<!ENTITY file SYSTEM "file:///private">]><doc>&file;</doc>')
        with self.assertRaises(KnowledgeBaseError) as caught:
            self.store.ingest_bytes(self.identity, "external.docx", output.getvalue())
        self.assertEqual(caught.exception.code, "unsafe_xml")

    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "pypdf must be installed by the parent")
    def test_no_text_pdf_does_not_create_a_success_document(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        output = io.BytesIO()
        writer.write(output)
        with self.assertRaises(KnowledgeBaseError) as caught:
            self.store.ingest_bytes(self.identity, "blank.pdf", output.getvalue())
        self.assertEqual(caught.exception.code, "no_text")
        self.assertEqual(self.store.list_documents(self.identity), [])
        self.assertEqual(list((self.root / "artifacts").iterdir()), [])

    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "pypdf must be installed by the parent")
    def test_real_pdf_retains_physical_page_and_full_chunk_quote(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 200 Td (" + b"pagetwotoken " * 20 + b") Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        output = io.BytesIO()
        writer.write(output)
        document = self.store.ingest_bytes(self.identity, "two-pages.pdf", output.getvalue(), "https://example.test/paper")
        hits = self.store.search(self.identity, "pagetwotoken")
        self.assertGreater(document["chunk_count"], 1)
        text = self.store.read_artifact(self.identity, document["text_artifact_id"]).decode("utf-8")
        for hit in hits:
            self.assertEqual(hit["citation"]["page"], 2)
            self.assertEqual(hit["citation"]["source"], "https://example.test/paper")
            self.assertEqual(hit["text"], text[hit["char_start"]:hit["char_end"]])
            self.assertEqual(self.store.get_chunk(self.identity, hit["chunk_id"])["text"], hit["text"])

    @unittest.skipUnless(importlib.util.find_spec("docx") and importlib.util.find_spec("openpyxl"), "office parsers are optional dependencies")
    def test_docx_and_xlsx_use_existing_parsers(self):
        from docx import Document
        from openpyxl import Workbook
        document = Document()
        document.add_heading("Document title", 1)
        document.add_paragraph("wordprocessingtoken")
        document.add_page_break()
        document.add_paragraph("secondpagetoken")
        output = io.BytesIO()
        document.save(output)
        self.store.ingest_bytes(self.identity, "document.docx", output.getvalue())
        self.assertEqual(self.store.search(self.identity, "secondpagetoken")[0]["page"], 2)
        workbook = Workbook()
        workbook.active.title = "Inventory"
        workbook.active.append(["spreadsheettoken", 12])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        self.store.ingest_bytes(self.identity, "workbook.xlsx", output.getvalue())
        self.assertEqual(self.store.search(self.identity, "spreadsheettoken")[0]["title"], "Inventory")

    def test_worker_roundtrip_and_timeout(self):
        store = KnowledgeBase(self.root / "worker", {**self.settings, "parse_in_worker": True})
        self.addCleanup(store.close)
        result = store.ingest_text(self.identity, "worker.txt", "isolatedparsertoken")
        self.assertTrue(result["ok"])
        self.assertTrue(store.search(self.identity, "isolatedparsertoken"))
        with patch.object(MODULE.subprocess, "run", side_effect=subprocess.TimeoutExpired("parser", 0.01)):
            with self.assertRaises(KnowledgeBaseError) as caught:
                store.ingest_text(self.identity, "timeout.txt", "different content")
        self.assertEqual(caught.exception.code, "parse_timeout")
        self.assertEqual(len(store.list_documents(self.identity)), 1)

    def test_url_batch_only_uses_supplied_fetcher(self):
        requested = []

        async def fetch(url):
            requested.append(url)
            if url.endswith("/empty"):
                return {"text": ""}
            return {"text": "externaldocumenttoken", "title": "Remote title"}

        results = asyncio.run(self.store.ingest_urls(self.identity, ["https://example.test/ok", "https://example.test/empty", "file:///etc/passwd"], fetch))
        self.assertTrue(results[0]["ok"])
        self.assertEqual(results[1]["error"], "no_text")
        self.assertEqual(results[2]["error"], "invalid_url")
        self.assertEqual(requested, ["https://example.test/ok", "https://example.test/empty"])
        hit = self.store.search(self.identity, "externaldocumenttoken")[0]
        self.assertEqual(hit["source"], "https://example.test/ok")
        self.assertEqual(hit["title"], "Remote title")


if __name__ == "__main__":
    unittest.main()