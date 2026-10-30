from __future__ import annotations

import asyncio
import base64
from contextlib import contextmanager, suppress
from dataclasses import asdict, dataclass, fields
import hashlib
from html.parser import HTMLParser
import inspect
import io
import json
import math
import os
from pathlib import Path, PurePosixPath
import re
import sqlite3
import stat
import subprocess
import sys
import tempfile
import threading
import time
from typing import Any, Mapping, TYPE_CHECKING
import unicodedata
from urllib.parse import urlsplit
import uuid
import zipfile

if TYPE_CHECKING:
    from .task_contracts import TaskIdentity


DEPENDENCIES = {".pdf": "pypdf", ".docx": "python-docx", ".xlsx": "openpyxl"}
SUPPORTED_EXTENSIONS = frozenset({".txt", ".md", ".markdown", ".html", ".htm", *DEPENDENCIES})
_ID = re.compile(r"[0-9a-f]{32}\Z")
_WORDS = re.compile(r"[a-z0-9_]+|[\u3400-\u9fff]+|[^\W_]+", re.UNICODE)


class KnowledgeBaseError(ValueError):
    def __init__(self, code: str, message: str = ""):
        super().__init__(message or code)
        self.code = code


@dataclass(frozen=True, slots=True)
class KnowledgeBaseSettings:
    max_upload_bytes: int = 20 * 1024 * 1024
    max_archive_bytes: int = 80 * 1024 * 1024
    max_archive_ratio: int = 100
    max_archive_entries: int = 2048
    max_text_chars: int = 2_000_000
    max_pages: int = 2000
    max_cells: int = 1_000_000
    chunk_size: int = 800
    chunk_overlap: int = 120
    max_chunks: int = 10000
    parse_timeout_seconds: float = 15.0
    parse_in_worker: bool = True
    semantic_enabled: bool = True
    local_bge_model: str = ""
    embedding_id: str = ""
    semantic_min_similarity: float = 0.15
    max_per_document: int = 2
    max_urls: int = 20

    @classmethod
    def from_value(cls, value: Any = None) -> KnowledgeBaseSettings:
        if isinstance(value, cls):
            value = asdict(value)
        if value is None:
            return cls()
        if not isinstance(value, Mapping):
            value = getattr(value, "catty_knowledge_base", value)
        if not isinstance(value, Mapping):
            value = {item.name: getattr(value, item.name) for item in fields(cls) if hasattr(value, item.name)}
        names = {item.name for item in fields(cls)}
        result = cls(**{key: item for key, item in value.items() if key in names})
        limits = {
            "max_upload_bytes": (1, 20 * 1024 * 1024),
            "max_archive_bytes": (1, 256 * 1024 * 1024),
            "max_archive_ratio": (1, 1000), "max_archive_entries": (1, 10000),
            "max_text_chars": (1, 10_000_000), "max_pages": (1, 10000),
            "max_cells": (1, 5_000_000), "chunk_size": (16, 4000),
            "max_chunks": (1, 50000), "max_per_document": (1, 100), "max_urls": (1, 100),
        }
        for name, (minimum, maximum) in limits.items():
            number = getattr(result, name)
            if type(number) is not int or not minimum <= number <= maximum:
                raise KnowledgeBaseError("invalid_settings", name)
        if type(result.chunk_overlap) is not int or not 0 <= result.chunk_overlap < result.chunk_size:
            raise KnowledgeBaseError("invalid_settings", "chunk_overlap")
        for name in ("parse_in_worker", "semantic_enabled"):
            if type(getattr(result, name)) is not bool:
                raise KnowledgeBaseError("invalid_settings", name)
        for name in ("local_bge_model", "embedding_id"):
            if not isinstance(getattr(result, name), str):
                raise KnowledgeBaseError("invalid_settings", name)
        for name, minimum, maximum in (("parse_timeout_seconds", 0.01, 300), ("semantic_min_similarity", -1, 1)):
            number = getattr(result, name)
            if isinstance(number, bool) or not isinstance(number, (int, float)) or not math.isfinite(number) or not minimum <= number <= maximum:
                raise KnowledgeBaseError("invalid_settings", name)
        return result


def _tokens(text: str, *, query: bool = False) -> list[str]:
    result = []
    for word in _WORDS.findall(unicodedata.normalize("NFKC", text).casefold()):
        if "\u3400" <= word[0] <= "\u9fff":
            if not query or len(word) == 1:
                result.extend(word)
            result.extend(word[index:index + 2] for index in range(len(word) - 1))
        else:
            result.append(word)
    return result


def _decode_text(content: bytes) -> str:
    encodings = ("utf-16",) if content.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "gb18030")
    for encoding in encodings:
        try:
            text = content.decode(encoding)
        except UnicodeError:
            continue
        if re.search(r"[\x00-\x08\x0b\x0e-\x1f]", text):
            raise KnowledgeBaseError("invalid_text")
        return text
    raise KnowledgeBaseError("invalid_encoding")


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.title: list[str] = []
        self.hidden = 0
        self.in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript", "template"}:
            self.hidden += 1
        if tag == "title":
            self.in_title = True
        if tag in {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "section"} and not self.hidden:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript", "template"}:
            self.hidden = max(0, self.hidden - 1)
        if tag == "title":
            self.in_title = False
        if tag in {"p", "div", "li", "tr", "h1", "h2", "h3", "section"} and not self.hidden:
            self.parts.append("\n")

    def handle_data(self, data):
        if self.in_title:
            self.title.append(data)
        elif not self.hidden:
            self.parts.append(data)


def _check_archive(content: bytes, settings: KnowledgeBaseSettings) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > settings.max_archive_entries:
                raise KnowledgeBaseError("archive_entries_exceeded")
            total = 0
            for entry in entries:
                name = entry.filename.replace("\\", "/")
                if (name.startswith("/") or ":" in name or ".." in PurePosixPath(name).parts
                        or any(ord(character) < 32 for character in name)
                        or stat.S_ISLNK(entry.external_attr >> 16) or entry.flag_bits & 1):
                    raise KnowledgeBaseError("unsafe_archive")
                total += entry.file_size
                if total > settings.max_archive_bytes:
                    raise KnowledgeBaseError("archive_size_exceeded")
                if entry.file_size > max(1, entry.compress_size) * settings.max_archive_ratio:
                    raise KnowledgeBaseError("archive_ratio_exceeded")
                if entry.is_dir():
                    continue
                actual = 0
                tail = b""
                with archive.open(entry) as stream:
                    while block := stream.read(65536):
                        actual += len(block)
                        if actual > entry.file_size or actual > settings.max_archive_bytes:
                            raise KnowledgeBaseError("archive_size_exceeded")
                        if name.lower().endswith((".xml", ".rels")):
                            scan = (tail + block).upper()
                            if b"<!DOCTYPE" in scan or b"<!ENTITY" in scan:
                                raise KnowledgeBaseError("unsafe_xml")
                            tail = scan[-16:]
    except (zipfile.BadZipFile, RuntimeError, NotImplementedError):
        raise KnowledgeBaseError("invalid_archive") from None


def _document_extension(content: bytes, settings: KnowledgeBaseSettings) -> str:
    if content.startswith(b"%PDF-"):
        return ".pdf"
    if content.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
        _check_archive(content, settings)
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = set(archive.namelist())
        if "[Content_Types].xml" in names:
            if "word/document.xml" in names:
                return ".docx"
            if "xl/workbook.xml" in names:
                return ".xlsx"
        raise KnowledgeBaseError("unsupported_format")
    if content.startswith((b"MZ", b"\x7fELF", b"\x89PNG", b"\xff\xd8\xff", b"GIF87a", b"GIF89a",
                           b"RIFF", b"OggS", b"fLaC", b"ID3", b"\x1f\x8b")) or content[4:8] == b"ftyp":
        raise KnowledgeBaseError("unsupported_format")
    try:
        text = _decode_text(content)
    except KnowledgeBaseError:
        raise KnowledgeBaseError("unsupported_format") from None
    if re.match(r"\s*(?:<!doctype\s+html\b|<html\b)", text, re.IGNORECASE):
        return ".html"
    if re.search(r"(?m)^#{1,6}\s+\S", text):
        return ".md"
    return ".txt"


def _parse_payload(content: bytes, extension: str, settings: KnowledgeBaseSettings) -> list[dict[str, Any]]:
    sections: list[dict[str, Any]] = []
    total_chars = 0

    def add(text, page=1, title=""):
        nonlocal total_chars
        text = str(text).replace("\r\n", "\n").replace("\r", "\n").strip()
        if not text:
            return
        total_chars += len(text) + 2
        if total_chars > settings.max_text_chars or len(sections) >= settings.max_chunks:
            raise KnowledgeBaseError("text_size_exceeded")
        sections.append({"text": text, "page": page, "title": str(title)[:500]})

    if extension in {".txt", ".md", ".markdown"}:
        pages = _decode_text(content).split("\f")
        if len(pages) > settings.max_pages:
            raise KnowledgeBaseError("page_limit_exceeded")
        for page_number, page_text in enumerate(pages, 1):
            title = ""
            paragraph: list[str] = []
            for line in page_text.splitlines():
                heading = re.match(r"^#{1,6}\s+(.+?)\s*#*\s*$", line) if extension != ".txt" else None
                if heading:
                    add("\n".join(paragraph), page_number, title)
                    paragraph = []
                    title = heading.group(1)
                paragraph.append(line)
            add("\n".join(paragraph), page_number, title)
    elif extension in {".html", ".htm"}:
        parser = _HTMLText()
        parser.feed(_decode_text(content))
        parser.close()
        add("".join(parser.parts), title="".join(parser.title).strip())
    elif extension == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise KnowledgeBaseError("dependency_missing", "pypdf") from None
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise KnowledgeBaseError("encrypted_document")
        if len(reader.pages) > settings.max_pages:
            raise KnowledgeBaseError("page_limit_exceeded")
        for page_number, page in enumerate(reader.pages, 1):
            add(page.extract_text() or "", page_number)
    elif extension == ".docx":
        _check_archive(content, settings)
        try:
            from docx import Document
            from docx.oxml.ns import qn
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError:
            raise KnowledgeBaseError("dependency_missing", "python-docx") from None
        document = Document(io.BytesIO(content))
        page_number = 1
        title = ""
        for element in document.element.body:
            if element.tag == qn("w:p"):
                paragraph = Paragraph(element, document)
                style_name = paragraph.style.name if paragraph.style is not None else ""
                if style_name.startswith(("Heading", "Title")):
                    title = paragraph.text
                add(paragraph.text, page_number, title)
                page_number += sum(1 for linebreak in element.iter(qn("w:br")) if linebreak.get(qn("w:type")) == "page")
                if page_number > settings.max_pages:
                    raise KnowledgeBaseError("page_limit_exceeded")
            elif element.tag == qn("w:tbl"):
                table = Table(element, document)
                add("\n".join("\t".join(cell.text for cell in row.cells) for row in table.rows), page_number, title)
    elif extension == ".xlsx":
        _check_archive(content, settings)
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise KnowledgeBaseError("dependency_missing", "openpyxl") from None
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        cell_count = 0
        try:
            if len(workbook.worksheets) > settings.max_pages:
                raise KnowledgeBaseError("page_limit_exceeded")
            for page_number, sheet in enumerate(workbook.worksheets, 1):
                if (sheet.max_row or 0) * (sheet.max_column or 0) > settings.max_cells:
                    raise KnowledgeBaseError("cell_limit_exceeded")
                rows = []
                sheet_chars = 0
                for row in sheet.iter_rows(values_only=True):
                    cell_count += len(row)
                    if cell_count > settings.max_cells:
                        raise KnowledgeBaseError("cell_limit_exceeded")
                    values = "\t".join("" if value is None else str(value) for value in row).rstrip()
                    sheet_chars += len(values) + 1
                    if total_chars + sheet_chars > settings.max_text_chars:
                        raise KnowledgeBaseError("text_size_exceeded")
                    if values:
                        rows.append(values)
                add("\n".join(rows), page_number, sheet.title)
        finally:
            workbook.close()
    else:
        raise KnowledgeBaseError("unsupported_format")
    if not sections:
        raise KnowledgeBaseError("no_text", "document contains no extractable text")
    return sections


def _worker_main() -> None:
    try:
        request = json.loads(sys.stdin.buffer.read(30 * 1024 * 1024))
        settings = KnowledgeBaseSettings.from_value(request["settings"])
        content = base64.b64decode(request["content"], validate=True)
        if len(content) > settings.max_upload_bytes:
            raise KnowledgeBaseError("upload_too_large")
        result = {"ok": True, "sections": _parse_payload(content, request["extension"], settings)}
    except Exception as exc:
        result = {"ok": False, "error": getattr(exc, "code", "parse_failed")}
    sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode("utf-8"))


class KnowledgeBase:
    """Private document storage. `path` and settings are trusted application configuration.

    Public IDs are opaque; no public method accepts a filesystem path. Embedding
    callbacks accept a list of strings and return equally many numeric vectors.
    """

    def __init__(self, path: str | Path, settings: Any = None, embedding_fn: Any = None):
        self.settings = KnowledgeBaseSettings.from_value(settings)
        self.root = Path(path).expanduser().resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self._artifacts = self.root / "artifacts"
        self._artifacts.mkdir(exist_ok=True)
        if self._artifacts.is_symlink() or self._artifacts.resolve().parent != self.root:
            raise KnowledgeBaseError("unsafe_storage")
        self.database_path = self.root / "knowledge.sqlite3"
        if self.database_path.is_symlink():
            raise KnowledgeBaseError("unsafe_storage")
        self._lock = threading.RLock()
        self._embedding_fn = embedding_fn if self.settings.semantic_enabled else None
        self._owns_embedding = False
        if self._embedding_fn is None and self.settings.semantic_enabled and self.settings.local_bge_model:
            from .rag_embedding_backend import LocalBGEEmbeddingFunction
            self._embedding_fn = LocalBGEEmbeddingFunction(model_name=self.settings.local_bge_model)
            self._owns_embedding = True
        backend = self._embedding_fn
        backend_name = getattr(backend, "name", None)
        if callable(backend_name):
            backend_name = backend_name()
        elif not isinstance(backend_name, str):
            backend_name = getattr(backend, "__module__", type(backend).__module__) + "." + getattr(backend, "__qualname__", type(backend).__qualname__)
        model_name = getattr(backend, "_model_name", self.settings.local_bge_model)
        self._space = self.settings.embedding_id or str(backend_name) + ":" + str(model_name)
        with self._db() as database:
            database.execute("PRAGMA journal_mode=WAL")
            database.executescript("""
                CREATE TABLE IF NOT EXISTS identities (
                    identity TEXT PRIMARY KEY, user_id TEXT NOT NULL, scope TEXT NOT NULL, persona TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS documents (
                    doc_id TEXT PRIMARY KEY, identity TEXT NOT NULL, name TEXT NOT NULL,
                    title TEXT NOT NULL, source TEXT NOT NULL, version INTEGER NOT NULL,
                    content_hash TEXT NOT NULL, artifact_id TEXT NOT NULL, text_artifact_id TEXT NOT NULL,
                    chunk_count INTEGER NOT NULL, semantic_pending INTEGER NOT NULL,
                    semantic_error TEXT NOT NULL, created_at REAL NOT NULL, updated_at REAL NOT NULL,
                    UNIQUE(identity, name)
                );
                CREATE INDEX IF NOT EXISTS documents_hash ON documents(identity, content_hash);
                CREATE TABLE IF NOT EXISTS artifacts (
                    artifact_id TEXT PRIMARY KEY, doc_id TEXT NOT NULL REFERENCES documents(doc_id) ON DELETE CASCADE,
                    identity TEXT NOT NULL, kind TEXT NOT NULL, size INTEGER NOT NULL
                );
                CREATE TABLE IF NOT EXISTS retired_artifacts (
                    artifact_id TEXT PRIMARY KEY, doc_id TEXT NOT NULL, identity TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS retired_identity ON retired_artifacts(identity,doc_id);
                CREATE TABLE IF NOT EXISTS chunks (
                    chunk_id TEXT PRIMARY KEY, doc_id TEXT NOT NULL REFERENCES documents(doc_id) ON DELETE CASCADE,
                    identity TEXT NOT NULL, version INTEGER NOT NULL, ordinal INTEGER NOT NULL,
                    page INTEGER NOT NULL, page_end INTEGER NOT NULL, title TEXT NOT NULL,
                    char_start INTEGER NOT NULL, char_end INTEGER NOT NULL, text TEXT NOT NULL, tokens TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS chunks_identity ON chunks(identity, doc_id, ordinal);
                CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(chunk_id UNINDEXED, identity UNINDEXED, tokens);
                CREATE TRIGGER IF NOT EXISTS chunks_insert AFTER INSERT ON chunks BEGIN
                    INSERT INTO chunks_fts(chunk_id, identity, tokens) VALUES(new.chunk_id, new.identity, new.tokens);
                END;
                CREATE TRIGGER IF NOT EXISTS chunks_delete AFTER DELETE ON chunks BEGIN
                    DELETE FROM chunks_fts WHERE chunk_id=old.chunk_id;
                END;
                CREATE TABLE IF NOT EXISTS vectors (
                    chunk_id TEXT PRIMARY KEY REFERENCES chunks(chunk_id) ON DELETE CASCADE,
                    space TEXT NOT NULL, dimensions INTEGER NOT NULL, vector TEXT NOT NULL
                );
                PRAGMA user_version=1;
            """)
            database.execute("""UPDATE documents SET semantic_pending=1, semantic_error='embedding_space_pending'
                WHERE chunk_count != (SELECT COUNT(*) FROM chunks JOIN vectors USING(chunk_id)
                WHERE chunks.doc_id=documents.doc_id AND vectors.space=?)""", (self._space,))

    @contextmanager
    def _db(self):
        with self._lock:
            database = sqlite3.connect(self.database_path, timeout=15)
            database.row_factory = sqlite3.Row
            database.execute("PRAGMA foreign_keys=ON")
            try:
                with database:
                    yield database
            finally:
                database.close()

    @staticmethod
    def _identity(identity: TaskIdentity) -> str:
        from .task_contracts import TaskIdentity
        if not isinstance(identity, TaskIdentity):
            raise TypeError("identity must be a verified TaskIdentity")
        values = [identity.user_id, identity.scope, identity.persona]
        return hashlib.sha256(json.dumps(values, ensure_ascii=False).encode("utf-8")).hexdigest()

    @staticmethod
    def _name(name: str) -> str:
        if (not isinstance(name, str) or not name.strip() or len(name) > 200
                or name != name.strip() or name in {".", ".."} or name.endswith((".", " "))
                or any(character in name for character in '/\\:<>|?*"')
                or any(ord(character) < 32 for character in name)):
            raise KnowledgeBaseError("invalid_name")
        if Path(name).suffix.lower() not in SUPPORTED_EXTENSIONS:
            raise KnowledgeBaseError("unsupported_format")
        return name

    def _artifact_path(self, artifact_id: str) -> Path:
        if not isinstance(artifact_id, str) or not _ID.fullmatch(artifact_id):
            raise KnowledgeBaseError("invalid_artifact")
        path = self._artifacts / artifact_id
        if path.is_symlink() or path.resolve().parent != self._artifacts.resolve():
            raise KnowledgeBaseError("unsafe_storage")
        return path

    def _parse(self, content: bytes, extension: str) -> list[dict[str, Any]]:
        if not self.settings.parse_in_worker:
            try:
                return _parse_payload(content, extension, self.settings)
            except KnowledgeBaseError:
                raise
            except Exception:
                raise KnowledgeBaseError("parse_failed") from None
        payload = json.dumps({"content": base64.b64encode(content).decode("ascii"), "extension": extension, "settings": asdict(self.settings)}).encode("utf-8")
        environment = {key: os.environ[key] for key in ("SYSTEMROOT", "WINDIR", "PATH", "TEMP", "TMP", "LANG") if key in os.environ}
        try:
            with tempfile.TemporaryDirectory(prefix=".parse-", dir=self.root) as directory:
                result = subprocess.run(
                    [sys.executable, "-I", "-B", "-X", "utf8", str(Path(__file__).resolve()), "--parse-worker"],
                    input=payload, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
                    cwd=directory, env=environment, timeout=self.settings.parse_timeout_seconds, check=False,
                )
            if result.returncode or len(result.stdout) > self.settings.max_text_chars * 6 + 2_000_000:
                raise KnowledgeBaseError("parse_failed")
            response = json.loads(result.stdout)
        except subprocess.TimeoutExpired:
            raise KnowledgeBaseError("parse_timeout") from None
        except (OSError, ValueError):
            raise KnowledgeBaseError("parse_failed") from None
        if not response.get("ok"):
            raise KnowledgeBaseError(response.get("error", "parse_failed"))
        return response["sections"]

    def _split(self, sections: list[dict[str, Any]]) -> tuple[str, list[dict[str, Any]]]:
        full_text = "\n\n".join(section["text"] for section in sections)
        if not full_text.strip():
            raise KnowledgeBaseError("no_text")
        if len(full_text) > self.settings.max_text_chars:
            raise KnowledgeBaseError("text_size_exceeded")
        chunks = []
        base_offset = 0
        for section in sections:
            text = section["text"]
            start = 0
            while start < len(text):
                end = min(len(text), start + self.settings.chunk_size)
                chunk = text[start:end]
                if chunk.strip():
                    chunks.append({"text": chunk, "page": section["page"], "page_end": section["page"],
                                   "title": section["title"], "char_start": base_offset + start, "char_end": base_offset + end})
                if len(chunks) > self.settings.max_chunks:
                    raise KnowledgeBaseError("chunk_limit_exceeded")
                if end == len(text):
                    break
                start = end - self.settings.chunk_overlap
            base_offset += len(text) + 2
        return full_text, chunks

    def _embed(self, texts: list[str], *, query: bool = False) -> list[list[float]]:
        if self._embedding_fn is None:
            raise KnowledgeBaseError("embedding_unconfigured")
        function = getattr(self._embedding_fn, "embed_query", self._embedding_fn) if query else self._embedding_fn
        output = []
        dimensions = None
        for offset in range(0, len(texts), 128):
            batch = texts[offset:offset + 128]
            vectors = function(batch)
            if hasattr(vectors, "tolist"):
                vectors = vectors.tolist()
            if not isinstance(vectors, (list, tuple)) or len(vectors) != len(batch):
                raise KnowledgeBaseError("invalid_embeddings")
            for vector in vectors:
                if hasattr(vector, "tolist"):
                    vector = vector.tolist()
                if not isinstance(vector, (list, tuple)) or not vector or len(vector) > 65536:
                    raise KnowledgeBaseError("invalid_embeddings")
                if any(isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) for value in vector):
                    raise KnowledgeBaseError("invalid_embeddings")
                norm = math.sqrt(sum(float(value) ** 2 for value in vector))
                if not math.isfinite(norm) or norm == 0 or dimensions not in (None, len(vector)):
                    raise KnowledgeBaseError("invalid_embeddings")
                dimensions = len(vector)
                output.append([float(value) / norm for value in vector])
        return output

    @staticmethod
    def _error(exc: Exception) -> str:
        return getattr(exc, "code", type(exc).__name__)

    @staticmethod
    def _document(row: sqlite3.Row, **extra) -> dict[str, Any]:
        result = dict(row)
        result.pop("identity", None)
        result["semantic_pending"] = bool(result["semantic_pending"])
        result["status"] = "semantic_pending" if result["semantic_pending"] else "indexed"
        return {"ok": True, **result, **extra}

    def ingest_bytes(self, identity: TaskIdentity, name: str, content: bytes, source: str = "") -> dict[str, Any]:
        try:
            self._name(name)
        except KnowledgeBaseError as exc:
            if exc.code != "unsupported_format" or Path(name).suffix.lower() not in {"", ".bin"}:
                raise
            if not isinstance(content, bytes):
                raise KnowledgeBaseError("invalid_content") from None
            if len(content) > self.settings.max_upload_bytes:
                raise KnowledgeBaseError("upload_too_large") from None
            name = self._name(name + _document_extension(content, self.settings))
        return self._ingest(identity, name, content, source)

    def ingest_text(self, identity: TaskIdentity, name: str, content: str, source: str = "") -> dict[str, Any]:
        if not isinstance(content, str):
            raise KnowledgeBaseError("invalid_text")
        if Path(self._name(name)).suffix.lower() not in {".txt", ".md", ".markdown", ".html", ".htm"}:
            raise KnowledgeBaseError("invalid_text_format")
        return self._ingest(identity, name, content.encode("utf-8"), source)

    def _ingest(self, identity, name, content, source, *, rebuild_row=None, title_override=""):
        identity_key = self._identity(identity)
        name = self._name(name)
        if not isinstance(source, str) or len(source) > 4096 or any(ord(character) < 32 for character in source):
            raise KnowledgeBaseError("invalid_source")
        if not isinstance(content, bytes):
            raise KnowledgeBaseError("invalid_content")
        if len(content) > self.settings.max_upload_bytes:
            raise KnowledgeBaseError("upload_too_large")
        digest = hashlib.sha256(content).hexdigest()
        with self._db() as database:
            duplicate = database.execute("SELECT * FROM documents WHERE identity=? AND content_hash=?", (identity_key, digest)).fetchone()
            if duplicate is not None and rebuild_row is None:
                return self._document(duplicate, duplicate=True)
        sections = self._parse(content, Path(name).suffix.lower())
        full_text, chunks = self._split(sections)
        embeddings = []
        semantic_error = ""
        try:
            embeddings = self._embed([chunk["text"] for chunk in chunks])
        except Exception as exc:
            semantic_error = self._error(exc)
        new_files: list[Path] = []
        old_files: list[Path] = []
        try:
            with self._db() as database:
                duplicate = database.execute("SELECT * FROM documents WHERE identity=? AND content_hash=?", (identity_key, digest)).fetchone()
                if duplicate is not None and rebuild_row is None:
                    return self._document(duplicate, duplicate=True)
                previous = database.execute("SELECT * FROM documents WHERE identity=? AND name=?", (identity_key, name)).fetchone()
                if rebuild_row is not None and (previous is None or previous["doc_id"] != rebuild_row["doc_id"]
                        or previous["version"] != rebuild_row["version"] or previous["artifact_id"] != rebuild_row["artifact_id"]):
                    raise KnowledgeBaseError("document_changed")
                doc_id = previous["doc_id"] if previous else uuid.uuid4().hex
                version = previous["version"] + (0 if rebuild_row is not None else 1) if previous else 1
                now = time.time()
                database.execute("INSERT OR IGNORE INTO identities VALUES(?,?,?,?)",
                    (identity_key, identity.user_id, identity.scope, identity.persona))
                artifact_id = previous["artifact_id"] if rebuild_row is not None else uuid.uuid4().hex
                text_artifact_id = uuid.uuid4().hex
                for new_id, data in ((artifact_id, content), (text_artifact_id, full_text.encode("utf-8"))):
                    if rebuild_row is not None and new_id == artifact_id:
                        continue
                    path = self._artifact_path(new_id)
                    new_files.append(path)
                    with path.open("xb") as stream:
                        stream.write(data)
                        stream.flush()
                        os.fsync(stream.fileno())
                if previous:
                    old_files = [self._artifact_path(previous["text_artifact_id"])]
                    if previous["artifact_id"] != artifact_id:
                        old_files.append(self._artifact_path(previous["artifact_id"]))
                    database.executemany("INSERT OR IGNORE INTO retired_artifacts VALUES(?,?,?)",
                        [(path.name, doc_id, identity_key) for path in old_files])
                    database.execute("DELETE FROM chunks WHERE doc_id=?", (doc_id,))
                    database.execute("DELETE FROM artifacts WHERE doc_id=?", (doc_id,))
                title = title_override or next((section["title"] for section in sections if section["title"]), name)
                database.execute("""INSERT INTO documents VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                    ON CONFLICT(doc_id) DO UPDATE SET title=excluded.title,source=excluded.source,version=excluded.version,
                    content_hash=excluded.content_hash,artifact_id=excluded.artifact_id,text_artifact_id=excluded.text_artifact_id,
                    chunk_count=excluded.chunk_count,semantic_pending=excluded.semantic_pending,
                    semantic_error=excluded.semantic_error,updated_at=excluded.updated_at""",
                    (doc_id, identity_key, name, title[:500], source, version, digest, artifact_id, text_artifact_id,
                     len(chunks), int(bool(semantic_error)), semantic_error, previous["created_at"] if previous else now, now))
                database.executemany("INSERT INTO artifacts VALUES(?,?,?,?,?)", [
                    (artifact_id, doc_id, identity_key, "original", len(content)),
                    (text_artifact_id, doc_id, identity_key, "text", len(full_text.encode("utf-8"))),
                ])
                for ordinal, chunk in enumerate(chunks):
                    chunk_id = hashlib.sha256(f"{doc_id}:{version}:{ordinal}:{chunk['text']}".encode("utf-8")).hexdigest()[:32]
                    database.execute("INSERT INTO chunks VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                        (chunk_id, doc_id, identity_key, version, ordinal, chunk["page"], chunk["page_end"], chunk["title"] or title,
                         chunk["char_start"], chunk["char_end"], chunk["text"], " ".join(_tokens(chunk["text"] + " " + chunk["title"]))))
                    if embeddings:
                        vector = embeddings[ordinal]
                        database.execute("INSERT INTO vectors VALUES(?,?,?,?)", (chunk_id, self._space, len(vector), json.dumps(vector)))
                row = database.execute("SELECT * FROM documents WHERE doc_id=?", (doc_id,)).fetchone()
        except Exception:
            for path in new_files:
                with suppress(OSError):
                    path.unlink(missing_ok=True)
            raise
        cleanup_pending = self._cleanup_artifacts(identity_key, doc_id)
        return self._document(row, duplicate=False, rebuilt=rebuild_row is not None, cleanup_pending=cleanup_pending)

    def _cleanup_artifacts(self, identity_key: str, doc_id: str) -> bool:
        pending = False
        with self._db() as database:
            rows = database.execute("SELECT artifact_id FROM retired_artifacts WHERE doc_id=? AND identity=?", (doc_id, identity_key)).fetchall()
            for row in rows:
                try:
                    self._artifact_path(row["artifact_id"]).unlink(missing_ok=True)
                except (OSError, KnowledgeBaseError):
                    pending = True
                else:
                    database.execute("DELETE FROM retired_artifacts WHERE artifact_id=? AND identity=?", (row["artifact_id"], identity_key))
        return pending

    def list_documents(self, identity: TaskIdentity) -> list[dict[str, Any]]:
        identity_key = self._identity(identity)
        with self._db() as database:
            rows = database.execute("SELECT * FROM documents WHERE identity=? ORDER BY updated_at DESC,doc_id", (identity_key,)).fetchall()
        return [self._document(row) for row in rows]

    def admin_list_documents(self) -> list[dict[str, Any]]:
        """Metadata-only entry point for an authenticated admin facade, never a model tool."""
        with self._db() as database:
            rows = database.execute("""SELECT documents.doc_id, identities.user_id, identities.scope, identities.persona,
                name,title,source,version,chunk_count,semantic_pending,semantic_error,created_at,updated_at
                FROM documents LEFT JOIN identities ON documents.identity=identities.identity
                ORDER BY updated_at DESC,documents.doc_id""").fetchall()
        return [{**dict(row), "semantic_pending": bool(row["semantic_pending"])} for row in rows]

    @staticmethod
    def _hit(row, **extra):
        result = dict(row)
        result.pop("identity", None)
        result.pop("tokens", None)
        result["semantic_pending"] = bool(result["semantic_pending"])
        result["citation"] = {key: result[key] for key in (
            "doc_id", "chunk_id", "version", "page", "page_end", "title", "source", "name",
            "artifact_id", "text_artifact_id", "char_start", "char_end",
        )}
        result["citation"]["ref"] = f"kb://{result['doc_id']}/v{result['version']}/{result['chunk_id']}"
        return {**result, **extra}

    @staticmethod
    def _chunk_query() -> str:
        return """SELECT chunks.*, documents.name, documents.source, documents.artifact_id,
            documents.text_artifact_id, documents.semantic_pending, documents.semantic_error
            FROM chunks JOIN documents ON documents.doc_id=chunks.doc_id"""

    def get_chunk(self, identity: TaskIdentity, chunk_id: str) -> dict[str, Any] | None:
        identity_key = self._identity(identity)
        with self._db() as database:
            row = database.execute(self._chunk_query() + " WHERE chunks.identity=? AND documents.identity=? AND chunk_id=?",
                (identity_key, identity_key, chunk_id)).fetchone()
        return self._hit(row) if row else None

    def read_artifact(self, identity: TaskIdentity, artifact_id: str) -> bytes | None:
        identity_key = self._identity(identity)
        with self._db() as database:
            row = database.execute("SELECT size FROM artifacts WHERE artifact_id=? AND identity=?", (artifact_id, identity_key)).fetchone()
            if row is None:
                return None
            path = self._artifact_path(artifact_id)
            try:
                with path.open("rb") as stream:
                    data = stream.read(row["size"] + 1)
            except FileNotFoundError:
                raise KnowledgeBaseError("artifact_missing") from None
            if len(data) != row["size"]:
                raise KnowledgeBaseError("artifact_corrupt")
            return data

    def search(self, identity: TaskIdentity, query: str, limit: int = 5) -> list[dict[str, Any]]:
        identity_key = self._identity(identity)
        if not isinstance(query, str) or len(query) > 4096 or type(limit) is not int or not 1 <= limit <= 100:
            raise KnowledgeBaseError("invalid_query")
        if not query.strip():
            return []
        terms = list(dict.fromkeys(_tokens(query, query=True)))[:128]
        candidates: dict[str, dict[str, Any]] = {}
        lexical_scores = {}
        semantic_scores = {}
        query_error = ""
        query_vector = None
        try:
            query_vector = self._embed([query], query=True)[0]
        except Exception as exc:
            query_error = self._error(exc)
        pool_size = min(800, max(40, limit * 8))
        with self._db() as database:
            if terms:
                expression = " OR ".join('"' + term.replace('"', '""') + '"' for term in terms)
                matches = database.execute("""SELECT chunk_id,bm25(chunks_fts) AS rank FROM chunks_fts
                    WHERE chunks_fts MATCH ? AND identity=? ORDER BY rank,chunk_id LIMIT ?""", (expression, identity_key, pool_size)).fetchall()
                lexical_scores = {row["chunk_id"]: index for index, row in enumerate(matches, 1)}
            if query_vector is not None:
                rows = database.execute("""SELECT vectors.chunk_id,vector FROM vectors JOIN chunks USING(chunk_id)
                    WHERE chunks.identity=? AND space=? AND dimensions=?""", (identity_key, self._space, len(query_vector)))
                for row in rows:
                    vector = json.loads(row["vector"])
                    similarity = sum(left * right for left, right in zip(query_vector, vector))
                    if similarity >= self.settings.semantic_min_similarity:
                        semantic_scores[row["chunk_id"]] = similarity
                semantic_scores = dict(sorted(semantic_scores.items(), key=lambda item: (-item[1], item[0]))[:pool_size])
            candidate_ids = set(lexical_scores) | set(semantic_scores)
            if candidate_ids:
                placeholders = ",".join("?" for _ in candidate_ids)
                rows = database.execute(self._chunk_query() + f" WHERE chunks.identity=? AND documents.identity=? AND chunk_id IN ({placeholders})",
                    (identity_key, identity_key, *candidate_ids))
                candidates = {row["chunk_id"]: dict(row) for row in rows}
        semantic_ranks = {chunk_id: rank for rank, chunk_id in enumerate(semantic_scores, 1)}
        scores = {}
        for chunk_id in candidates:
            score = 1 / (60 + lexical_scores[chunk_id]) if chunk_id in lexical_scores else 0
            if chunk_id in semantic_ranks:
                score += 1 / (60 + semantic_ranks[chunk_id]) + 0.01 * semantic_scores[chunk_id]
            scores[chunk_id] = score
        selected = []
        doc_counts: dict[str, int] = {}
        texts: set[str] = set()
        while candidates and len(selected) < limit:
            eligible = [chunk_id for chunk_id, row in candidates.items() if doc_counts.get(row["doc_id"], 0) < self.settings.max_per_document]
            if not eligible:
                eligible = list(candidates)
            chunk_id = max(eligible, key=lambda item: (
                scores[item] / (1 + 0.6 * doc_counts.get(candidates[item]["doc_id"], 0)) * (0.1 if candidates[item]["text"] in texts else 1), item))
            row = candidates.pop(chunk_id)
            doc_counts[row["doc_id"]] = doc_counts.get(row["doc_id"], 0) + 1
            texts.add(row["text"])
            selected.append(self._hit(row, score=scores[chunk_id], lexical_rank=lexical_scores.get(chunk_id),
                semantic_score=semantic_scores.get(chunk_id), search_mode="hybrid" if semantic_scores else "fts",
                semantic_pending=bool(row["semantic_pending"] or query_error), semantic_error=query_error or row["semantic_error"]))
        return selected

    def retry_pending(self, identity: TaskIdentity, limit: int = 10) -> dict[str, Any]:
        identity_key = self._identity(identity)
        if type(limit) is not int or not 1 <= limit <= 100:
            raise KnowledgeBaseError("invalid_limit")
        with self._db() as database:
            documents = database.execute("SELECT * FROM documents WHERE identity=? AND semantic_pending=1 ORDER BY updated_at LIMIT ?", (identity_key, limit)).fetchall()
        indexed = 0
        failed = 0
        for document in documents:
            with self._db() as database:
                chunks = database.execute("SELECT chunk_id,text FROM chunks WHERE doc_id=? AND identity=? ORDER BY ordinal", (document["doc_id"], identity_key)).fetchall()
            vectors = []
            error = ""
            try:
                if not chunks:
                    raise KnowledgeBaseError("chunks_missing")
                vectors = self._embed([row["text"] for row in chunks])
            except Exception as exc:
                error = self._error(exc)
            with self._db() as database:
                current = database.execute("SELECT version,text_artifact_id FROM documents WHERE doc_id=? AND identity=?", (document["doc_id"], identity_key)).fetchone()
                if current is None or current["version"] != document["version"] or current["text_artifact_id"] != document["text_artifact_id"]:
                    continue
                if not error:
                    for chunk, vector in zip(chunks, vectors):
                        database.execute("INSERT OR REPLACE INTO vectors VALUES(?,?,?,?)", (chunk["chunk_id"], self._space, len(vector), json.dumps(vector)))
                    indexed += 1
                else:
                    failed += 1
                database.execute("UPDATE documents SET semantic_pending=?,semantic_error=? WHERE doc_id=?", (int(bool(error)), error, document["doc_id"]))
        with self._db() as database:
            pending = database.execute("SELECT COUNT(*) FROM documents WHERE identity=? AND semantic_pending=1", (identity_key,)).fetchone()[0]
        return {"ok": failed == 0, "indexed": indexed, "failed": failed, "pending": pending}

    def rebuild(self, identity: TaskIdentity, doc_id: str) -> dict[str, Any]:
        identity_key = self._identity(identity)
        with self._db() as database:
            row = database.execute("SELECT * FROM documents WHERE doc_id=? AND identity=?", (doc_id, identity_key)).fetchone()
        if row is None:
            return {"ok": False, "error": "not_found", "doc_id": doc_id}
        content = self.read_artifact(identity, row["artifact_id"])
        if content is None or hashlib.sha256(content).hexdigest() != row["content_hash"]:
            raise KnowledgeBaseError("artifact_corrupt")
        return self._ingest(identity, row["name"], content, row["source"], rebuild_row=row)

    def delete_document(self, identity: TaskIdentity, doc_id: str) -> dict[str, Any]:
        identity_key = self._identity(identity)
        with self._db() as database:
            document = database.execute("SELECT doc_id FROM documents WHERE doc_id=? AND identity=?", (doc_id, identity_key)).fetchone()
            retired = database.execute("SELECT 1 FROM retired_artifacts WHERE doc_id=? AND identity=? LIMIT 1", (doc_id, identity_key)).fetchone()
            if document is None and retired is None:
                return {"ok": False, "deleted": False, "error": "not_found", "doc_id": doc_id}
            database.execute("""INSERT OR IGNORE INTO retired_artifacts SELECT artifact_id,doc_id,identity
                FROM artifacts WHERE doc_id=? AND identity=?""", (doc_id, identity_key))
            database.execute("DELETE FROM documents WHERE doc_id=? AND identity=?", (doc_id, identity_key))
        cleanup_pending = self._cleanup_artifacts(identity_key, doc_id)
        return {"ok": not cleanup_pending, "deleted": True, "doc_id": doc_id, "cleanup_pending": cleanup_pending}

    async def ingest_urls(self, identity: TaskIdentity, urls: list[str], public_fetch) -> list[dict[str, Any]]:
        """Use a caller-supplied SSRF-checked fetcher returning text or {text, title}."""
        self._identity(identity)
        if not isinstance(urls, (list, tuple)) or len(urls) > self.settings.max_urls or not callable(public_fetch):
            raise KnowledgeBaseError("invalid_urls")
        results = []
        for url in urls:
            try:
                if not isinstance(url, str) or len(url) > 4096:
                    raise KnowledgeBaseError("invalid_url")
                parts = urlsplit(url)
                if parts.scheme not in {"http", "https"} or not parts.hostname or parts.username or parts.password:
                    raise KnowledgeBaseError("invalid_url")
                if inspect.iscoroutinefunction(public_fetch):
                    response = await public_fetch(url)
                else:
                    response = await asyncio.to_thread(public_fetch, url)
                if inspect.isawaitable(response):
                    response = await response
                text = response.get("text", "") if isinstance(response, Mapping) else response
                title = response.get("title", "") if isinstance(response, Mapping) else ""
                if not isinstance(text, str) or not text.strip() or not isinstance(title, str):
                    raise KnowledgeBaseError("no_text")
                name = "url-" + hashlib.sha256(url.encode("utf-8")).hexdigest()[:16] + ".txt"
                result = await asyncio.to_thread(self._ingest, identity, name, text.encode("utf-8"), url, title_override=title)
            except Exception as exc:
                result = {"ok": False, "error": self._error(exc)}
            results.append({**result, "url": url})
        return results

    def close(self) -> None:
        if self._owns_embedding:
            self._embedding_fn.close()


if __name__ == "__main__" and sys.argv[1:] == ["--parse-worker"]:
    _worker_main()
