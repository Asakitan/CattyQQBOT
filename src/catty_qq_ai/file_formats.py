"""文件交付编码：真实 DOCX/XLSX/PDF，其他后缀明确作为 UTF-8 原文。"""
from __future__ import annotations

import csv
import io
import os
import tempfile
import threading
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape


DOCUMENT_FORMATS = frozenset({"docx", "xlsx", "pdf"})
# 这是不支持的编码器提示，不是文件后缀白名单；filename 任意后缀仍可写原文。
_BINARY_FORMATS = frozenset({
    "binary", "bin", "doc", "xls", "ppt", "pptx", "odt", "ods", "odp",
    "zip", "rar", "7z", "gz", "tar", "png", "jpg", "jpeg", "gif", "webp",
    "bmp", "ico", "mp3", "mp4", "wav", "avi", "exe", "dll", "so",
})
_MEDIA_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "text": "text/plain; charset=utf-8",
}
_FONT_LOCK = threading.Lock()


def resolve_output_format(filename: str, requested: str = "") -> tuple[str, str]:
    """显式格式优先；真实文档修正扩展名；text/raw 不更改任意后缀。"""
    requested = requested.strip().lower().lstrip(".")
    if requested in _BINARY_FORMATS:
        raise ValueError(
            f"不支持 {requested} 二进制编码；请用 catty_python_run 生成真实产物。"
            "若只需这个后缀的 UTF-8 原文，请明确使用 format=text 或 raw。"
        )
    suffix = Path(filename).suffix.lstrip(".").lower()
    if requested in DOCUMENT_FORMATS:
        return str(Path(filename).with_suffix("." + requested)), requested
    if requested:
        # md/csv/json/自定义文本格式不做内容转换，也不假称已生成二进制。
        return filename, "text"
    return filename, suffix if suffix in DOCUMENT_FORMATS else "text"


def _docx_bytes(content: str) -> bytes:
    from docx import Document

    document = Document()
    for line in content.split("\n"):
        document.add_paragraph(line.rstrip("\r"))
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _xlsx_bytes(content: str) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "内容"
    # 固定 CSV 规则：逗号分列、双引号引用、支持引号内换行；普通文本是一列。
    # 保留所有字段为文字（前导零、=开头原文等），不猜类型、不执行公式。
    try:
        for row_index, row in enumerate(csv.reader(io.StringIO(content, newline=""), strict=True), 1):
            if row_index > 1_048_576 or len(row) > 16_384:
                raise ValueError("内容超出 XLSX 行列上限")
            for column_index, value in enumerate(row, 1):
                if len(value) > 32_767:
                    raise ValueError("XLSX 单元格超过 32767 字符，请拆成多行或改用文本文件")
                cell = worksheet.cell(row_index, column_index, value)
                cell.data_type = "s"
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()
    finally:
        workbook.close()


def _pdf_bytes(content: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    # ReportLab 自带 CID 字体度量，不下载字体；阅读器负责显示标准中文 CID 字体。
    font_name = "STSong-Light"
    with _FONT_LOCK:
        try:
            pdfmetrics.getFont(font_name)
        except KeyError:
            pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    style = ParagraphStyle("CattyDocument", fontName=font_name, fontSize=11,
                           leading=17, wordWrap="CJK", spaceAfter=4)
    story = []
    for line in content.split("\n"):
        line = line.rstrip("\r")
        if line:
            story.append(Paragraph(escape(line).replace("\t", "&#160;" * 4), style))
        else:
            story.append(Spacer(1, style.leading))
    stream = io.BytesIO()
    document = SimpleDocTemplate(stream, pagesize=A4, leftMargin=40, rightMargin=40,
                                 topMargin=40, bottomMargin=40)
    document.build(story)
    return stream.getvalue()


def write_document_atomic(path: Path, content: str, output_format: str) -> dict[str, Any]:
    """同步编码 + fsync + 同目录 replace；调用方必须放到 asyncio.to_thread。"""
    if output_format == "docx":
        payload = _docx_bytes(content)
    elif output_format == "xlsx":
        payload = _xlsx_bytes(content)
    elif output_format == "pdf":
        payload = _pdf_bytes(content)
    elif output_format == "text":
        payload = content.encode("utf-8")
    else:
        raise ValueError(f"未知输出编码: {output_format}")

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        # 同目录临时文件，失败不暴露半个文档，也不覆盖其他并发请求的临时文件。
        with tempfile.NamedTemporaryFile(mode="wb", dir=path.parent, suffix=".tmp", delete=False) as stream:
            temporary = Path(stream.name)
            stream.write(payload)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)

    metadata: dict[str, Any] = {
        "format": output_format,
        "bytes": len(payload),
        "media_type": _MEDIA_TYPES[output_format],
        "content_mode": "text/raw" if output_format == "text" else "document",
    }
    if output_format == "text":
        metadata["encoding"] = "utf-8"
        metadata["format_note"] = (
            "实际内容为 UTF-8 原文；后缀不代表已编码为该二进制格式。"
            "需要其他真实二进制文件时请用 catty_python_run。"
        )
    elif output_format == "xlsx":
        metadata["format_note"] = "正文按逗号 CSV 解析，双引号可包住逗号/换行；所有单元格保留为文本。"
    elif output_format == "pdf":
        metadata["format_note"] = "使用离线 STSong-Light 中文 CID 字体；未联网下载字体。"
    return metadata